#!/usr/bin/env python3
"""Bouw de zoekmodule-data en de statische pagina's onder /vinden/.

  python3 scripts/build.py            -> valideert data/ en schrijft vinden/data/vinden.json,
                                         vinden/<organisatie>/, vinden/werksoort/<x>/[<provincie>/],
                                         vinden/sitemap-vinden.xml
  python3 scripts/build.py --export   -> ook "2026 WKN Vinden data.xlsx" in ~/Downloads (voor Google Sheets)

Alle pagina's krijgen noindex tot het launchbesluit (zet INDEXEREN = True bij launch).
"""
import argparse, csv, html, json, re, shutil, sys, unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

INDEXEREN = False
MIN_PROVINCIE = 3
BASIS = "https://werkkrachtnederland.nl"
ROOT = Path(__file__).resolve().parent.parent
UIT = ROOT / "vinden"
NL_BBOX = (50.7, 53.7, 3.2, 7.3)
UITV = ("eigen_bedrijf", "op_locatie", "beide")
e = html.escape


def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def lees_csv(naam):
    return list(csv.DictReader(open(ROOT / "data" / naam, encoding="utf-8")))


# ---------- laden en valideren ----------
def laad():
    tax = json.load(open(ROOT / "data/taxonomie.json", encoding="utf-8"))
    orgs = lees_csv("organisaties.csv")
    vest = lees_csv("vestigingen.csv")
    aanbod = lees_csv("aanbod.csv") if (ROOT / "data/aanbod.csv").exists() else []
    fouten = []
    hg_ids = {h["id"] for h in tax["hoofdgroepen"]}
    ws = {w["id"]: w for w in tax["werksoorten"]}
    for w in tax["werksoorten"]:
        if w["hoofdgroep"] not in hg_ids:
            fouten.append(f"werksoort {w['id']}: onbekende hoofdgroep {w['hoofdgroep']}")
        for veld in ("naam_en", "naam_es"):
            if not w.get(veld):
                fouten.append(f"werksoort {w['id']}: {veld} ontbreekt (drietaligheid)")
        if w.get("uitvoeringsvorm_standaard") not in UITV:
            fouten.append(f"werksoort {w['id']}: ongeldige uitvoeringsvorm_standaard")
    ob = {o["id"]: o for o in orgs}
    if len({o["slug"] for o in orgs}) != len(orgs):
        fouten.append("dubbele organisatie-slugs")
    vb = defaultdict(list)
    for v in vest:
        if v["organisatie_id"] not in ob:
            fouten.append(f"vestiging {v['id']}: onbekende organisatie")
            continue
        try:
            la, lo = float(v["lat"]), float(v["lon"])
            if not (NL_BBOX[0] <= la <= NL_BBOX[1] and NL_BBOX[2] <= lo <= NL_BBOX[3]):
                fouten.append(f"vestiging {v['id']}: coördinaten buiten Nederland")
        except ValueError:
            fouten.append(f"vestiging {v['id']}: geen geldige coördinaten")
            continue
        vb[v["organisatie_id"]].append(v)
    ab = defaultdict(list)
    gezien = set()
    for a in aanbod:
        if a["organisatie_id"] not in ob:
            fouten.append(f"aanbod: onbekende organisatie {a['organisatie_id']}")
        elif a["werksoort_id"] not in ws:
            fouten.append(f"aanbod: onbekende werksoort {a['werksoort_id']}")
        elif a.get("uitvoeringsvorm") and a["uitvoeringsvorm"] not in UITV:
            fouten.append(f"aanbod: ongeldige uitvoeringsvorm {a['uitvoeringsvorm']}")
        elif (a["organisatie_id"], a["werksoort_id"]) in gezien:
            fouten.append(f"aanbod: dubbel {a['organisatie_id']}/{a['werksoort_id']}")
        else:
            gezien.add((a["organisatie_id"], a["werksoort_id"]))
            ab[a["organisatie_id"]].append(a)
    for o in orgs:
        if not vb[o["id"]]:
            fouten.append(f"organisatie {o['slug']}: geen vestiging")
    if fouten:
        print("VALIDATIE MISLUKT:\n  " + "\n  ".join(fouten))
        sys.exit(1)
    return tax, orgs, vb, ab, ws


# ---------- HTML-bouwstenen ----------
BEELDMERK = ('<svg viewBox="0 0 72 72" width="34" height="34" aria-hidden="true">'
             '<rect x="2" y="56" width="22" height="9" rx="2.5" fill="#27AE60"/>'
             '<rect x="14" y="42" width="22" height="9" rx="2.5" fill="#A9DFBF"/>'
             '<rect x="26" y="28" width="22" height="9" rx="2.5" fill="#FFFFFF"/></svg>')


def pagina(titel, omschrijving, pad, inhoud, versie, jsonld=None):
    robots = "index,follow" if INDEXEREN else "noindex,nofollow"
    ld = ""
    if jsonld:
        ld = '<script type="application/ld+json">' + json.dumps(jsonld, ensure_ascii=False).replace("</", "<\\/") + "</script>"
    return f"""<!DOCTYPE html>
<html lang="nl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(titel)}</title>
<meta name="description" content="{e(omschrijving)}">
<meta name="robots" content="{robots}">
<link rel="canonical" href="{BASIS}{pad}">
<meta property="og:title" content="{e(titel)}">
<meta property="og:description" content="{e(omschrijving)}">
<meta property="og:url" content="{BASIS}{pad}">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Syne:wght@700;800&family=Red+Hat+Display:wght@700;900&family=Red+Hat+Text:wght@400;500;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="/vinden/assets/vinden.css?v={versie}">
{ld}
</head>
<body>
<header class="kop">
  <a class="merk" href="/">{BEELDMERK}<span>WerkKracht Nederland</span></a>
  <nav class="kop-nav"><a href="/vinden/" data-t="Zoek een werkontwikkelbedrijf">Zoek een werkontwikkelbedrijf</a>
  <span class="taal" role="group" aria-label="Taal"><button type="button" data-lang="nl">NL</button><button type="button" data-lang="en">EN</button><button type="button" data-lang="es">ES</button></span></nav>
</header>
<div class="strip"></div>
<main>{inhoud}</main>
<footer class="voet">
  <p class="payoff" data-t="Elke stap telt.">Elke stap telt.</p>
  <p data-t="Gegevens afgeleid van de openbare websites van de werkontwikkelbedrijven. Klopt er iets niet? Laat het ons weten.">Gegevens afgeleid van de openbare websites van de werkontwikkelbedrijven. Klopt er iets niet? Laat het ons weten.</p>
  <p><a href="/#contact" data-t="Neem contact op met WerkKracht Nederland">Neem contact op met WerkKracht Nederland</a></p>
</footer>
<script src="/vinden/assets/taal.js?v={versie}"></script>
</body>
</html>
"""


def wsnaam(w):
    return f'<span class="ws" data-en="{e(w["naam_en"])}" data-es="{e(w["naam_es"])}">{e(w["naam"])}</span>'


def uitvnaam(tax, u):
    nl, en, es = tax["uitvoeringsvormen"][u]
    return f'<span class="ws" data-en="{e(en)}" data-es="{e(es)}">{e(nl)}</span>'


def profiel(o, vest, aanbod, tax, ws, versie):
    pad = f"/vinden/{o['slug']}/"
    per_hg = defaultdict(list)
    for a in aanbod:
        w = ws[a["werksoort_id"]]
        per_hg[w["hoofdgroep"]].append((w, a.get("uitvoeringsvorm") or w["uitvoeringsvorm_standaard"]))
    hgn = {h["id"]: h for h in tax["hoofdgroepen"]}
    blokken = ""
    for hid in [h["id"] for h in tax["hoofdgroepen"] if h["id"] in per_hg]:
        h = hgn[hid]
        regels = "".join(f'<li><a href="/vinden/werksoort/{w["slug"]}/">{wsnaam(w)}</a><small>{uitvnaam(tax, u)}</small></li>'
                         for w, u in sorted(per_hg[hid], key=lambda x: x[0]["naam"]))
        blokken += f'<section class="hg"><h3>{wsnaam(h)}</h3><ul class="wslijst">{regels}</ul></section>'
    if not blokken:
        blokken = '<p class="leeg" data-t="De werksoorten van deze organisatie zijn nog niet bekend.">De werksoorten van deze organisatie zijn nog niet bekend.</p>'
    def vregel(v):
        delen = [v["adres"], " ".join(x for x in (v["postcode"], v["plaats"]) if x)]
        return "<li>" + e(", ".join(x for x in delen if x)) + "</li>"
    vlijst = "".join(vregel(v) for v in vest)
    zoek = "/vinden/?w=" + ",".join(sorted({ws[a["werksoort_id"]]["slug"] for a in aanbod}))
    inhoud = f"""
<article class="profiel">
  <p class="kruimel"><a href="/vinden/" data-t="Zoek een werkontwikkelbedrijf">Zoek een werkontwikkelbedrijf</a></p>
  <h1>{e(o['naam'])}</h1>
  <p class="plaats">{e(o['plaats'])}, {e(o['provincie'])}</p>
  {f'<p class="intro">{e(o["beschrijving"])}</p>' if o.get("beschrijving") else ''}
  <p class="acties"><a class="knop" href="{e(o['website'])}" rel="noopener" target="_blank" data-t="Bezoek de website">Bezoek de website</a>
  <a class="knop knop-licht" href="/#contact" data-t="Vraag via WerkKracht Nederland">Vraag via WerkKracht Nederland</a></p>
  <div class="profiel-raster">
    <div><h2 data-t="Werksoorten">Werksoorten</h2>{blokken}
    <p class="noot" data-t="De uitvoeringsvorm is indicatief. Vraag de actuele mogelijkheden na bij de organisatie.">De uitvoeringsvorm is indicatief. Vraag de actuele mogelijkheden na bij de organisatie.</p></div>
    <aside><h2 data-t="Vestigingen">Vestigingen</h2><ul class="vlijst">{vlijst}</ul>
    {f'<p><a href="{zoek}" data-t="Vergelijk met andere werkontwikkelbedrijven">Vergelijk met andere werkontwikkelbedrijven</a></p>' if aanbod else ''}</aside>
  </div>
</article>"""
    ld = {"@context": "https://schema.org", "@type": "Organization", "name": o["naam"], "url": o["website"],
          "address": [{"@type": "PostalAddress", "addressLocality": v["plaats"], "addressRegion": v["provincie"],
                       "postalCode": v["postcode"] or None, "streetAddress": v["adres"] or None, "addressCountry": "NL"} for v in vest],
          "knowsAbout": [ws[a["werksoort_id"]]["naam"] for a in aanbod]}
    titel = f"{o['naam']} in {o['plaats']}: werksoorten en contact"
    oms = f"{o['naam']} is een werkontwikkelbedrijf in {o['plaats']}. Bekijk werksoorten, vestigingen en contactmogelijkheden."
    return pad, pagina(titel, oms, pad, inhoud, versie, ld)


def lijstpagina(w, orgs_hier, tax, versie, provincie=None, provincies=None):
    pslug = slugify(provincie) if provincie else None
    pad = f"/vinden/werksoort/{w['slug']}/" + (f"{pslug}/" if provincie else "")
    per_prov = defaultdict(list)
    for o, u in orgs_hier:
        per_prov[o["provincie"]].append((o, u))
    blokken = ""
    for prov in sorted(per_prov):
        items = "".join(f'<li><a href="/vinden/{o["slug"]}/">{e(o["naam"])}</a><span>{e(o["plaats"])}</span><small>{uitvnaam(tax, u)}</small></li>'
                        for o, u in sorted(per_prov[prov], key=lambda x: x[0]["naam"]))
        kop = "" if provincie else f'<h2>{e(prov)}</h2>'
        blokken += f'<section>{kop}<ul class="orglijst">{items}</ul></section>'
    provlinks = ""
    if provincies and not provincie:
        provlinks = '<p class="provlinks">' + " ".join(f'<a href="{pad}{slugify(p)}/">{e(p)}</a>' for p in provincies) + "</p>"
    waar = f" in {provincie}" if provincie else ""
    titel = f"{w['naam']} uitbesteden{waar} aan werkontwikkelbedrijven"
    intro = w.get("intro") or f"Werkontwikkelbedrijven{waar} voeren {w['naam'].lower()} uit voor opdrachtgevers. Kies een partij bij u in de buurt."
    inhoud = f"""
<article class="lijst">
  <p class="kruimel"><a href="/vinden/" data-t="Zoek een werkontwikkelbedrijf">Zoek een werkontwikkelbedrijf</a>{f' / <a href="/vinden/werksoort/{w["slug"]}/">{wsnaam(w)}</a>' if provincie else ''}</p>
  <h1>{e(titel)}</h1>
  <p class="intro">{e(intro)}</p>
  <p class="acties"><a class="knop" href="/vinden/?w={w['slug']}" data-t="Zoek op afstand van uw locatie">Zoek op afstand van uw locatie</a></p>
  <p class="aantal">{len(orgs_hier)} <span data-t="werkontwikkelbedrijven">werkontwikkelbedrijven</span></p>
  {provlinks}{blokken}
</article>"""
    ld = {"@context": "https://schema.org", "@type": "ItemList", "name": titel,
          "itemListElement": [{"@type": "ListItem", "position": i + 1, "url": f"{BASIS}/vinden/{o['slug']}/", "name": o["naam"]}
                              for i, (o, _) in enumerate(sorted(orgs_hier, key=lambda x: x[0]["naam"]))]}
    oms = f"{len(orgs_hier)} werkontwikkelbedrijven{waar} voor {w['naam'].lower()}. Vergelijk op afstand en werksoort."
    return pad, pagina(titel, oms, pad, inhoud, versie, ld)


def schrijf(pad, inhoud):
    doel = ROOT / pad.lstrip("/") / "index.html"
    doel.parent.mkdir(parents=True, exist_ok=True)
    doel.write_text(inhoud, encoding="utf-8")


def export_xlsx(tax, orgs, vb, ab):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)

    def blad(naam, kop, rijen):
        ws_ = wb.create_sheet(naam)
        ws_.append(kop)
        for c in ws_[1]:
            c.font = Font(bold=True)
        for r in rijen:
            ws_.append(r)
        ws_.freeze_panes = "A2"

    blad("organisaties", list(orgs[0].keys()), [list(o.values()) for o in orgs])
    alle_v = [v for lst in vb.values() for v in lst]
    blad("vestigingen", list(alle_v[0].keys()), [list(v.values()) for v in alle_v])
    blad("hoofdgroepen", ["id", "naam", "naam_en", "naam_es", "slug"], [[h[k] for k in ("id", "naam", "naam_en", "naam_es", "slug")] for h in tax["hoofdgroepen"]])
    blad("werksoorten", ["id", "hoofdgroep", "naam", "naam_en", "naam_es", "slug", "termen", "uitvoeringsvorm_standaard", "intro"],
         [[w["id"], w["hoofdgroep"], w["naam"], w["naam_en"], w["naam_es"], w["slug"], ", ".join(w["termen"]), w["uitvoeringsvorm_standaard"], w.get("intro", "")] for w in tax["werksoorten"]])
    alle_a = [a for lst in ab.values() for a in lst]
    if alle_a:
        blad("aanbod", list(alle_a[0].keys()), [list(a.values()) for a in alle_a])
    blad("zoeklog", ["datum", "werksoorten", "postcode4", "straal_km", "alles_vereist", "aantal_resultaten", "taal"], [])
    doel = Path.home() / "Downloads" / "2026 WKN Vinden data.xlsx"
    doel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(doel)
    print(f"Export: {doel}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--export", action="store_true")
    a = ap.parse_args()
    tax, orgs, vb, ab, ws = laad()
    versie = datetime.now(timezone.utc).strftime("%Y%m%d%H%M")

    # opruimen van eerder gegenereerde pagina's (index.html, assets en data blijven staan)
    for p in UIT.iterdir() if UIT.exists() else []:
        if p.is_dir() and p.name not in ("assets", "data"):
            shutil.rmtree(p)

    # vinden.json
    data = {
        "versie": versie,
        "uitvoeringsvormen": tax["uitvoeringsvormen"],
        "hoofdgroepen": [{k: h[k] for k in ("id", "naam", "naam_en", "naam_es", "slug")} for h in tax["hoofdgroepen"]],
        "werksoorten": [{"id": w["id"], "hg": w["hoofdgroep"], "naam": w["naam"], "naam_en": w["naam_en"], "naam_es": w["naam_es"],
                         "slug": w["slug"], "termen": w["termen"]} for w in tax["werksoorten"]],
        "organisaties": [],
    }
    for o in orgs:
        if not ab[o["id"]]:
            continue
        data["organisaties"].append({
            "id": o["id"], "naam": o["naam"], "slug": o["slug"], "plaats": o["plaats"], "provincie": o["provincie"], "website": o["website"],
            "vestigingen": [[round(float(v["lat"]), 5), round(float(v["lon"]), 5), v["plaats"]] for v in vb[o["id"]]],
            "aanbod": [[x["werksoort_id"], x.get("uitvoeringsvorm") or ws[x["werksoort_id"]]["uitvoeringsvorm_standaard"]] for x in ab[o["id"]]],
        })
    (UIT / "data").mkdir(parents=True, exist_ok=True)
    (UIT / "data/vinden.json").write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    paden = ["/vinden/"]
    for o in orgs:
        pad, h = profiel(o, vb[o["id"]], ab[o["id"]], tax, ws, versie)
        schrijf(pad, h)
        paden.append(pad)

    per_ws = defaultdict(list)
    for o in orgs:
        for x in ab[o["id"]]:
            per_ws[x["werksoort_id"]].append((o, x.get("uitvoeringsvorm") or ws[x["werksoort_id"]]["uitvoeringsvorm_standaard"]))
    n_prov = 0
    for wid, lijst in per_ws.items():
        w = ws[wid]
        tel = defaultdict(list)
        for o, u in lijst:
            tel[o["provincie"]].append((o, u))
        provs = sorted(p for p, l in tel.items() if len(l) >= MIN_PROVINCIE)
        pad, h = lijstpagina(w, lijst, tax, versie, provincies=provs)
        schrijf(pad, h)
        paden.append(pad)
        for p in provs:
            pad, h = lijstpagina(w, tel[p], tax, versie, provincie=p)
            schrijf(pad, h)
            paden.append(pad)
            n_prov += 1

    vandaag = datetime.now(timezone.utc).date().isoformat()
    sm = '<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
    sm += "".join(f"  <url><loc>{BASIS}{p}</loc><lastmod>{vandaag}</lastmod></url>\n" for p in paden)
    sm += "</urlset>\n"
    (UIT / "sitemap-vinden.xml").write_text(sm, encoding="utf-8")

    # versienummer in index.html bijwerken (cache-busting van css/js/json)
    idx = UIT / "index.html"
    if idx.exists():
        t = idx.read_text(encoding="utf-8")
        idx.write_text(re.sub(r"\?v=\d{12}", f"?v={versie}", t), encoding="utf-8")

    zichtbaar = len(data["organisaties"])
    print(f"OK: {zichtbaar} organisaties zoekbaar, {len(orgs)} profielen, {len(per_ws)} werksoortpagina's, "
          f"{n_prov} provinciepagina's, {len(paden)} URL's in sitemap. Indexeren: {INDEXEREN}")
    if a.export:
        export_xlsx(tax, orgs, vb, ab)


if __name__ == "__main__":
    main()
